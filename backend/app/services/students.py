"""Student registry: edits, consent changes and roster imports (CSV / XLSX)."""
from __future__ import annotations

import csv
import io
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.audit import ledger
from app.core.errors import Conflict, NotFound, ValidationFailed
from app.core.hashing import sha256_hex
from app.db.base import utcnow
from app.domain.normalize import (
    clean_name,
    is_valid_roll,
    normalize_branch,
    normalize_roll,
    parse_backlogs,
    parse_cgpa,
    parse_consent,
    parse_year,
)
from app.models.entities import ConsentRecord, RosterImport, Student, User
from app.models.enums import ActorType, ConsentScope, ConsentStatus, ImportStatus
from app.services.approvals import revalidate_open_approvals
from app.services.consent import current_consent

# Target fields and the header spellings we auto-map. Officers can override.
FIELD_SYNONYMS: dict[str, list[str]] = {
    "roll_number": ["roll_number", "roll no", "roll no.", "rollno", "roll", "registration number", "reg no", "usn", "enrollment no"],
    "full_name": ["full_name", "name", "student name", "full name", "candidate name"],
    "email": ["email", "email id", "e-mail", "mail"],
    "phone": ["phone", "mobile", "contact", "phone number"],
    "branch": ["branch", "department", "dept", "discipline", "stream"],
    "cgpa": ["cgpa", "cpi", "gpa", "aggregate cgpa"],
    "active_backlogs": ["active_backlogs", "backlogs", "active backlogs", "arrears", "current backlogs"],
    "graduation_year": ["graduation_year", "batch", "passing year", "year of passing", "graduation year"],
    "consent": ["consent", "consent_share", "data sharing consent", "consent to share", "share consent"],
}
REQUIRED_FIELDS = ("roll_number", "full_name")


def _actor(user: User | None) -> str:
    return f"user:{user.id}" if user else "system"


def get_student(db: Session, student_id: int) -> Student:
    s = db.get(Student, student_id)
    if s is None:
        raise NotFound(f"Student {student_id} not found")
    return s


# --------------------------------------------------------------------------- import

def read_table(filename: str, content: bytes) -> tuple[list[str], list[list[Any]]]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise ValidationFailed("Only .csv and .xlsx files are supported")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValidationFailed("File has no rows")
    headers = [str(h).strip() for h in rows[0]]
    return headers, rows[1:]


def auto_mapping(headers: list[str]) -> dict[str, str | None]:
    lowered = {h.lower().strip(): h for h in headers}
    mapping: dict[str, str | None] = {}
    for field, synonyms in FIELD_SYNONYMS.items():
        mapping[field] = next((lowered[s] for s in synonyms if s in lowered), None)
    return mapping


def _normalise_row(raw: dict[str, Any]) -> tuple[dict[str, Any], list[dict]]:
    issues: list[dict] = []
    roll = normalize_roll(raw.get("roll_number"))
    if not roll:
        issues.append({"field": "roll_number", "code": "MISSING", "message": "Roll number is required", "severity": "ERROR"})
    elif not is_valid_roll(roll):
        issues.append({"field": "roll_number", "code": "INVALID", "message": f"Roll number '{roll}' has an invalid format", "severity": "ERROR"})
    name = clean_name(raw.get("full_name"))
    if not name:
        issues.append({"field": "full_name", "code": "MISSING", "message": "Name is required", "severity": "ERROR"})
    branch_raw = clean_name(raw.get("branch"))
    branch = normalize_branch(branch_raw)
    if branch_raw and branch is None:
        issues.append({"field": "branch", "code": "UNRECOGNISED", "message": f"Branch '{branch_raw}' is not recognised", "severity": "WARNING"})
    elif not branch_raw:
        issues.append({"field": "branch", "code": "MISSING", "message": "Branch is missing", "severity": "WARNING"})
    cgpa = parse_cgpa(raw.get("cgpa"))
    backlogs = parse_backlogs(raw.get("active_backlogs"))
    year = parse_year(raw.get("graduation_year"))
    consent = parse_consent(raw.get("consent"))
    for parsed in (cgpa, backlogs, year, consent):
        issues.extend(i.as_dict() for i in parsed.issues)
    email = clean_name(raw.get("email"))
    if email and "@" not in email:
        issues.append({"field": "email", "code": "INVALID", "message": "Email looks invalid", "severity": "WARNING"})
    normalised = {
        "roll_number": roll, "full_name": name, "email": email, "phone": clean_name(raw.get("phone")),
        "branch": branch, "branch_raw": branch_raw, "cgpa": cgpa.value, "active_backlogs": backlogs.value,
        "graduation_year": year.value, "consent": consent.value,
    }
    return normalised, issues


def preview_import(db: Session, filename: str, content: bytes, mapping: dict[str, str | None] | None, user: User) -> RosterImport:
    headers, rows = read_table(filename, content)
    mapping = mapping or auto_mapping(headers)
    for field in REQUIRED_FIELDS:
        if not mapping.get(field):
            raise ValidationFailed(f"Column mapping for '{field}' is required", details={"headers": headers, "mapping": mapping})
    unknown = [h for h in mapping.values() if h and h not in headers]
    if unknown:
        raise ValidationFailed(f"Mapped columns not in file: {unknown}", details={"headers": headers})
    index = {h: i for i, h in enumerate(headers)}

    existing = {s.roll_number: s for s in db.scalars(select(Student))}
    seen: dict[str, int] = {}
    preview: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for line_no, row in enumerate(rows, start=2):
        raw = {f: (row[index[col]] if col and index[col] < len(row) else None) for f, col in mapping.items()}
        normalised, issues = _normalise_row(raw)
        roll = normalised["roll_number"]
        if roll and roll in seen:
            issues.append({"field": "roll_number", "code": "DUPLICATE_IN_FILE",
                           "message": f"Roll number {roll} also appears on line {seen[roll]}", "severity": "ERROR"})
        elif roll:
            seen[roll] = line_no
        blocking = any(i["severity"] == "ERROR" for i in issues)
        action = "REJECT"
        if not blocking:
            current = existing.get(roll)
            if current is None:
                action = "CREATE"
            else:
                changed = any(getattr(current, k) != normalised[k] for k in
                              ("full_name", "email", "phone", "branch", "branch_raw", "cgpa", "active_backlogs", "graduation_year"))
                action = "UPDATE" if changed else "UNCHANGED"
                if current.full_name.lower() != (normalised["full_name"] or "").lower():
                    issues.append({"field": "full_name", "code": "NAME_CHANGED",
                                   "message": f"Roll {roll} exists as '{current.full_name}'; name will change",
                                   "severity": "WARNING"})
        preview.append({"line": line_no, "action": action, "data": normalised, "issues": issues})
        for i in issues:
            errors.append({"line": line_no, "roll_number": roll, **i})

    imp = RosterImport(
        filename=filename, file_sha256=sha256_hex(content), status=ImportStatus.PREVIEWED,
        column_mapping=mapping, total_rows=len(rows),
        valid_rows=sum(p["action"] != "REJECT" for p in preview),
        error_rows=sum(p["action"] == "REJECT" for p in preview),
        preview_rows=preview, errors=errors, created_by_id=user.id,
    )
    db.add(imp)
    db.flush()
    ledger.record(db, actor_type=ActorType.USER, actor_id=_actor(user), action="IMPORT_PREVIEWED",
                  entity_type="roster_import", entity_id=imp.id,
                  details={"filename": filename, "sha256": imp.file_sha256, "rows": imp.total_rows,
                           "valid": imp.valid_rows, "rejected": imp.error_rows})
    return imp


def commit_import(db: Session, import_id: int, user: User) -> RosterImport:
    imp = db.get(RosterImport, import_id)
    if imp is None:
        raise NotFound("Import not found")
    if imp.status != ImportStatus.PREVIEWED:
        raise Conflict(f"Import is already {imp.status}")
    existing = {s.roll_number: s for s in db.scalars(select(Student))}
    created = updated = unchanged = 0
    touched: list[Student] = []
    consent_changes: list[tuple[Student, bool | None]] = []
    for row in imp.preview_rows:
        if row["action"] == "REJECT":
            continue
        d = row["data"]
        warnings = [i for i in row["issues"] if i["severity"] == "WARNING"]
        s = existing.get(d["roll_number"])
        if s is None:
            s = Student(roll_number=d["roll_number"], full_name=d["full_name"], email=d["email"], phone=d["phone"],
                        branch=d["branch"], branch_raw=d["branch_raw"], cgpa=d["cgpa"],
                        active_backlogs=d["active_backlogs"], graduation_year=d["graduation_year"],
                        data_issues=warnings, last_import_id=imp.id)
            db.add(s)
            created += 1
        else:
            changed = False
            for k in ("full_name", "email", "phone", "branch", "branch_raw", "cgpa", "active_backlogs", "graduation_year"):
                if getattr(s, k) != d[k]:
                    setattr(s, k, d[k])
                    changed = True
            s.data_issues = warnings
            s.last_import_id = imp.id
            if changed:
                s.record_version += 1
                updated += 1
            else:
                unchanged += 1
        touched.append(s)
        consent_changes.append((s, d["consent"]))
    db.flush()

    scope = ConsentScope.SHARE_WITH_EMPLOYERS
    current = current_consent(db, [s.id for s, _ in consent_changes], scope)
    for s, value in consent_changes:
        desired = ConsentStatus.GRANTED if value is True else ConsentStatus.REVOKED
        now = current.get(s.id)
        if value is None and now is None:
            continue  # nothing recorded and nothing stated: remains "not consented"
        if now is None or now["status"] != desired:
            db.add(ConsentRecord(student_id=s.id, scope=scope, status=desired, source=f"import:{imp.id}",
                                 evidence_ref=f"{imp.filename} sha256:{imp.file_sha256[:12]}", recorded_by_id=user.id))
    imp.status = ImportStatus.COMMITTED
    imp.committed_at = utcnow()
    imp.created_count, imp.updated_count, imp.unchanged_count = created, updated, unchanged
    db.flush()
    ledger.record(db, actor_type=ActorType.USER, actor_id=_actor(user), action="IMPORT_COMMITTED",
                  entity_type="roster_import", entity_id=imp.id,
                  details={"created": created, "updated": updated, "unchanged": unchanged, "rejected": imp.error_rows})
    revalidate_open_approvals(db, f"Roster import #{imp.id} changed student records or consent")
    return imp


def errors_csv(imp: RosterImport) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["line", "roll_number", "field", "severity", "code", "message"])
    for e in imp.errors:
        w.writerow([e.get("line"), e.get("roll_number"), e.get("field"), e.get("severity"), e.get("code"), e.get("message")])
    return buf.getvalue()


# --------------------------------------------------------------------------- edits

EDITABLE = ("full_name", "email", "phone", "branch", "cgpa", "active_backlogs", "graduation_year")


def update_student(db: Session, student: Student, data: dict[str, Any], user: User) -> dict[str, Any]:
    changes: dict[str, Any] = {}
    for key in EDITABLE:
        if key not in data:
            continue
        value = data[key]
        if key == "branch":
            code = normalize_branch(value)
            if value and code is None:
                raise ValidationFailed(f"Unrecognised branch '{value}'")
            value = code
            if value != student.branch:
                student.branch_raw = data[key]
        elif key == "cgpa" and value is not None:
            parsed = parse_cgpa(value)
            if parsed.value is None:
                raise ValidationFailed(parsed.issues[0].message)
            value = parsed.value
        elif key == "active_backlogs" and value is not None:
            parsed = parse_backlogs(value)
            if parsed.value is None:
                raise ValidationFailed(parsed.issues[0].message)
            value = parsed.value
        elif key == "full_name":
            value = clean_name(value)
            if not value:
                raise ValidationFailed("Name cannot be empty")
        if getattr(student, key) != value:
            changes[key] = {"from": getattr(student, key), "to": value}
            setattr(student, key, value)
    if not changes:
        return {"changes": {}, "invalidated_proposals": []}
    student.record_version += 1
    student.data_issues = [i for i in (student.data_issues or []) if i.get("field") not in changes]
    db.flush()
    ledger.record(db, actor_type=ActorType.USER, actor_id=_actor(user), action="STUDENT_UPDATED",
                  entity_type="student", entity_id=student.id,
                  details={"roll_number": student.roll_number, "changes": changes, "record_version": student.record_version})
    invalidated = revalidate_open_approvals(db, f"Student {student.roll_number} record changed")
    return {"changes": changes, "invalidated_proposals": invalidated}


def set_consent(db: Session, student: Student, granted: bool, evidence_ref: str | None, user: User | None,
                source: str = "manual") -> dict[str, Any]:
    status = ConsentStatus.GRANTED if granted else ConsentStatus.REVOKED
    rec = ConsentRecord(student_id=student.id, scope=ConsentScope.SHARE_WITH_EMPLOYERS, status=status,
                        source=source, evidence_ref=evidence_ref, recorded_by_id=user.id if user else None)
    db.add(rec)
    db.flush()
    ledger.record(db, actor_type=ActorType.USER if user else ActorType.SYSTEM, actor_id=_actor(user),
                  action="CONSENT_GRANTED" if granted else "CONSENT_REVOKED", entity_type="student",
                  entity_id=student.id, details={"roll_number": student.roll_number, "scope": rec.scope,
                                                 "evidence_ref": evidence_ref})
    invalidated = revalidate_open_approvals(db, f"Consent changed for {student.roll_number}")
    return {"consent_record_id": rec.id, "invalidated_proposals": invalidated}
